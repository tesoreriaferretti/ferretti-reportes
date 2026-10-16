from datetime import datetime
import io
import re
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import pandas as pd
import streamlit as st


def clean_op(val):
  if pd.isna(val) or val is None:
    return ""
  val_str = str(val).strip().replace(".0", "")
  val_clean = re.sub(r"[^a-zA-Z0-9]", "", val_str)
  return val_clean.lstrip("0") if val_clean.lstrip("0") else val_clean


def format_op_display(val):
  if pd.isna(val) or val is None:
    return ""
  s = str(val).strip()
  return s[:-2] if s.endswith(".0") else s


def safe_float(val):
  if pd.isna(val) or val is None:
    return 0.0
  try:
    return float(str(val).replace(",", "").strip())
  except Exception:
    return 0.0


def smart_parse_date(val):
  if pd.isna(val) or val is None:
    return pd.NaT
  if isinstance(val, (datetime, pd.Timestamp)):
    return pd.to_datetime(val)
  s = str(val).strip()
  if re.match(r"^\d{4}[-/]\d{1,2}[-/]\d{1,2}", s):
    return pd.to_datetime(s, format="mixed", errors="coerce")
  return pd.to_datetime(s, dayfirst=True, errors="coerce")


def is_white_fill(cell):
  if not cell or not cell.fill or not cell.fill.fill_type:
    return True
  fg = cell.fill.fgColor
  if not fg:
    return True
  if fg.type == "theme" and fg.theme == 0 and abs(fg.tint or 0.0) < 0.01:
    return True
  if fg.type == "rgb" and fg.rgb in ["FFFFFFFF", "00000000", None]:
    return True
  return False


def clean_comercio_nombre(raw_comercio, default_sheet_name=""):
  val = str(raw_comercio or "").strip().upper()
  if "MIRAFLORES" in val:
    return "Miraflores"
  elif "MOLINA" in val:
    return "La Molina"
  elif "LINK" in val or "Link" in default_sheet_name:
    return "Link"
  elif "POS" in default_sheet_name:
    return str(raw_comercio or "POS").strip()
  return str(raw_comercio or "").strip()


def render():
  st.subheader("Reporte Ejecutivo: Pendientes / No Identificados")

  uploaded_files = st.file_uploader(
      "Sube los archivos necesarios (Hoja Bancos Auditada, Extracto BCP, FP"
      " Tarjetas, Pendientes Anterior):",
      type=["xlsx"],
      accept_multiple_files=True,
      key="uploader_no_identificados",
  )

  if st.button("Generar Reporte de Pendientes", type="primary"):
    if not uploaded_files:
      st.error("Debes subir los archivos requeridos.")
      return

    file_dict = {f.name: f for f in uploaded_files}
    bcp_file = next(
        (
            f
            for n, f in file_dict.items()
            if "bcp" in n.lower() or "movimiento" in n.lower()
        ),
        None,
    )
    tarj_file = next(
        (f for n, f in file_dict.items() if "tarjeta" in n.lower()), None
    )
    pend_ant_file = next(
        (f for n, f in file_dict.items() if "pendiente" in n.lower()), None
    )
    bancos_auditado_file = next(
        (
            f
            for n, f in file_dict.items()
            if "banco" in n.lower()
            and f != bcp_file
            and "tarjeta" not in n.lower()
            and "pendiente" not in n.lower()
        ),
        None,
    )

    if not (bancos_auditado_file and tarj_file):
      st.error(
          "Es obligatorio subir la hoja auditada de Bancos y el archivo de FP"
          " Tarjetas."
      )
      return

    with st.spinner("Extrayendo depósitos y movimientos no identificados..."):
      wb_bancos = openpyxl.load_workbook(bancos_auditado_file, data_only=True)
      ws_b = wb_bancos.active

      mapa_usuario_bcp = {}
      if bcp_file:
        try:
          df_raw_bcp = pd.read_excel(bcp_file, header=None)
          h_idx = 0
          for idx, row in df_raw_bcp.head(15).iterrows():
            row_str = " ".join([str(v).lower() for v in row.values])
            if "fecha" in row_str and "monto" in row_str:
              h_idx = idx
              break
          bcp_file.seek(0)
          df_bcp_ref = pd.read_excel(bcp_file, header=h_idx, dtype=str)
          df_bcp_ref.columns = [str(c).strip() for c in df_bcp_ref.columns]
          op_col = [c for c in df_bcp_ref.columns if "operac" in c.lower()]
          usr_col = [c for c in df_bcp_ref.columns if "usuario" in c.lower()]
          if op_col and usr_col:
            df_bcp_ref["op_clean"] = df_bcp_ref[op_col[0]].apply(clean_op)
            mapa_usuario_bcp = (
                df_bcp_ref.dropna(subset=[usr_col[0]])
                .set_index("op_clean")[usr_col[0]]
                .to_dict()
            )
        except Exception as e:
          st.warning(f"Aviso lectura BCP: {e}")

      unmatched_soles = []
      unmatched_dolares = []

      for r in range(2, ws_b.max_row + 1):
        f_b = ws_b.cell(row=r, column=12).value
        desc_b = ws_b.cell(row=r, column=13).value
        m_b = ws_b.cell(row=r, column=14).value
        op_b = ws_b.cell(row=r, column=15).value
        mon_b = ws_b.cell(row=r, column=16).value
        obs_b = ws_b.cell(row=r, column=17).value

        if m_b is None and f_b is None:
          continue

        c_monto = ws_b.cell(row=r, column=14)
        is_colored = False
        if c_monto.fill and c_monto.fill.fill_type:
          fg = c_monto.fill.fgColor
          theme = getattr(fg, "theme", None)
          rgb = getattr(fg, "rgb", None)
          if theme in [9, 5, 4, 6] or (
              rgb and any(c in str(rgb) for c in ["FFFF00", "F8CB", "DDEB"])
          ):
            is_colored = True

        obs_str = str(obs_b).strip().lower() if obs_b is not None else ""
        is_special = any(
            k in obs_str
            for k in [
                "duplicado",
                "cambio de moneda",
                "no realizada",
                "extorno",
                "ok tarjeta",
                "f010",
                "b010",
            ]
        )

        if not is_colored and not is_special and obs_str in ["", "none", "nan"]:
          op_cl = clean_op(op_b)
          entry = {
              "Fecha": f_b,
              "Descripción operación": str(desc_b or ""),
              "Monto": safe_float(m_b),
              "Nº operación": format_op_display(op_b),
              "Usuario": mapa_usuario_bcp.get(op_cl, ""),
              "Detalle de la Operación - Rspta BCP": "",
              "RUC / DNI": "",
              "op_clean": op_cl,
          }
          if str(mon_b).strip().lower() == "soles":
            unmatched_soles.append(entry)
          else:
            unmatched_dolares.append(entry)

      df_soles_nuevos = pd.DataFrame(unmatched_soles)
      df_dolares_nuevos = pd.DataFrame(unmatched_dolares)

      def consolidar_banco_historico(df_nuevos, sheet_name, pend_file):
        if not pend_file:
          return df_nuevos
        try:
          wb_p = openpyxl.load_workbook(pend_file, data_only=True)
          if sheet_name not in wb_p.sheetnames:
            return df_nuevos
          ws_p = wb_p[sheet_name]

          headers = {}
          for c in range(1, ws_p.max_column + 1):
            val = ws_p.cell(row=4, column=c).value
            if val:
              headers[str(val).strip()] = c

          col_f = headers.get("Fecha")
          col_desc = headers.get("Descripción operación")
          col_m = headers.get("Monto")
          col_op = headers.get("Nº operación")
          col_usr = headers.get("Usuario")
          col_det = headers.get("Detalle de la Operación - Rspta BCP")
          col_doc = headers.get("RUC / DNI")

          hist_rows = []
          for r in range(5, ws_p.max_row + 1):
            m_val = ws_p.cell(row=r, column=col_m).value if col_m else None
            op_val = ws_p.cell(row=r, column=col_op).value if col_op else None
            if (
                m_val is not None
                and str(m_val).strip() not in ["", "None", "Monto"]
            ):
              hist_rows.append({
                  "Fecha": (
                      ws_p.cell(row=r, column=col_f).value if col_f else ""
                  ),
                  "Descripción operación": str(
                      ws_p.cell(row=r, column=col_desc).value or ""
                  ),
                  "Monto": safe_float(m_val),
                  "Nº operación": format_op_display(op_val),
                  "Usuario": (
                      str(ws_p.cell(row=r, column=col_usr).value or "")
                      if col_usr
                      else ""
                  ),
                  "Detalle de la Operación - Rspta BCP": (
                      str(ws_p.cell(row=r, column=col_det).value or "")
                      if col_det
                      else ""
                  ),
                  "RUC / DNI": (
                      str(ws_p.cell(row=r, column=col_doc).value or "")
                      if col_doc
                      else ""
                  ),
                  "op_clean": clean_op(op_val),
              })
          df_hist = pd.DataFrame(hist_rows)
          if df_hist.empty:
            return df_nuevos

          map_det = (
              df_hist[df_hist["Detalle de la Operación - Rspta BCP"] != ""]
              .set_index("op_clean")["Detalle de la Operación - Rspta BCP"]
              .to_dict()
          )
          map_doc = (
              df_hist[df_hist["RUC / DNI"] != ""]
              .set_index("op_clean")["RUC / DNI"]
              .to_dict()
          )

          if not df_nuevos.empty:
            df_nuevos["Detalle de la Operación - Rspta BCP"] = (
                df_nuevos["op_clean"].map(map_det).fillna("")
            )
            df_nuevos["RUC / DNI"] = (
                df_nuevos["op_clean"].map(map_doc).fillna("")
            )
            ops_actuales = set(df_nuevos["op_clean"])
          else:
            ops_actuales = set()

          hist_pendientes = df_hist[
              ~df_hist["op_clean"].isin(ops_actuales)
          ].copy()
          return pd.concat(
              [df_nuevos, hist_pendientes], ignore_index=True
          ).drop_duplicates(subset=["op_clean"])
        except Exception:
          return df_nuevos

      soles_final = consolidar_banco_historico(
          df_soles_nuevos, "FP Soles", pend_ant_file
      )
      dolares_final = consolidar_banco_historico(
          df_dolares_nuevos, "FP Dolares", pend_ant_file
      )

      tarjetas_blanco_list = []
      wb_tarj = openpyxl.load_workbook(tarj_file, data_only=True)
      for s in [
          "Culqi Link MN",
          "Culqi Link ME",
          "Culqi POS MN",
          "Culqi POS ME",
      ]:
        if s in wb_tarj.sheetnames:
          ws_t = wb_tarj[s]
          headers = {
              str(ws_t.cell(row=4, column=c).value).strip(): c
              for c in range(1, ws_t.max_column + 1)
          }

          col_id = headers.get("ID Transaccion")
          col_est = headers.get("Estado")
          col_monto = headers.get("Venta Final")
          col_fecha = headers.get("Fecha de la transaccion")
          col_cli = (
              headers.get("Nombre tarjetahabiente")
              or headers.get("Nombre y Apellido")
              or headers.get("Cliente")
          )

          if not col_cli:
            for h, c in headers.items():
              if any(k in h.lower() for k in ["tarjetahabiente", "nombre"]):
                col_cli = c
                break

          moneda_std = (
              "Dolares" if ("ME" in s or "usd" in str(s).lower()) else "Soles"
          )

          for r in range(5, ws_t.max_row + 1):
            est_val = (
                str(ws_t.cell(row=r, column=col_est).value).strip().lower()
                if col_est
                else ""
            )
            if est_val in ["aprobada", "abonada"]:
              c_ref = ws_t.cell(row=r, column=col_id or 4)
              c_est = ws_t.cell(row=r, column=col_est)
              if is_white_fill(c_ref) and is_white_fill(c_est):
                raw_id = (
                    ws_t.cell(row=r, column=col_id).value if col_id else ""
                )
                val_monto = (
                    ws_t.cell(row=r, column=col_monto).value if col_monto else 0.0
                )
                val_fecha = (
                    ws_t.cell(row=r, column=col_fecha).value if col_fecha else ""
                )
                val_cli = (
                    ws_t.cell(row=r, column=col_cli).value
                    if col_cli
                    else "POR IDENTIFICAR"
                )

                raw_comercio = ws_t.cell(row=r, column=2).value
                comercio_final = clean_comercio_nombre(
                    raw_comercio, default_sheet_name=s
                )

                tarjetas_blanco_list.append({
                    "RUC": "20603632479",
                    "Cod. Comercio": comercio_final,
                    "Moneda": moneda_std,
                    "Fecha operación": val_fecha,
                    "ID Operación": format_op_display(raw_id),
                    "Monto": safe_float(val_monto),
                    "Cliente": str(val_cli).strip(),
                    "op_clean": clean_op(raw_id),
                })

      df_tarj_nuevas = pd.DataFrame(tarjetas_blanco_list)

      def consolidar_tarjetas_historico(df_nuevas, pend_file):
        if not pend_file:
          return df_nuevas
        try:
          wb_p = openpyxl.load_workbook(pend_file, data_only=True)
          if "TARJETAS" not in wb_p.sheetnames:
            return df_nuevas
          ws_p = wb_p["TARJETAS"]

          headers = {
              str(ws_p.cell(row=4, column=c).value).strip(): c
              for c in range(1, ws_p.max_column + 1)
              if ws_p.cell(row=4, column=c).value
          }
          col_ruc = headers.get("RUC")
          col_com = headers.get("Cod. Comercio")
          col_mon = headers.get("Moneda")
          col_f = headers.get("Fecha operación")
          col_id = headers.get("ID Operación")
          col_m = headers.get("Monto")
          col_cli = headers.get("Cliente")

          if not col_id or not col_m:
            return df_nuevas

          hist_rows = []
          for r in range(5, ws_p.max_row + 1):
            m_val = ws_p.cell(row=r, column=col_m).value
            id_val = ws_p.cell(row=r, column=col_id).value
            if (
                m_val is not None
                and str(m_val).strip() not in ["", "None", "Monto"]
            ):
              raw_com = (
                  ws_p.cell(row=r, column=col_com).value if col_com else ""
              )
              hist_rows.append({
                  "RUC": (
                      str(ws_p.cell(row=r, column=col_ruc).value or "20603632479")
                      if col_ruc
                      else "20603632479"
                  ),
                  "Cod. Comercio": clean_comercio_nombre(raw_com),
                  "Moneda": (
                      str(ws_p.cell(row=r, column=col_mon).value or "").strip()
                      if col_mon
                      else ""
                  ),
                  "Fecha operación": (
                      ws_p.cell(row=r, column=col_f).value if col_f else ""
                  ),
                  "ID Operación": format_op_display(id_val),
                  "Monto": safe_float(m_val),
                  "Cliente": (
                      str(ws_p.cell(row=r, column=col_cli).value or "").strip()
                      if col_cli
                      else ""
                  ),
                  "op_clean": clean_op(id_val),
              })

          df_hist = pd.DataFrame(hist_rows)
          if df_hist.empty:
            return df_nuevas
          ops_actuales = (
              set(df_nuevas["op_clean"]) if not df_nuevas.empty else set()
          )
          hist_pendientes = df_hist[
              ~df_hist["op_clean"].isin(ops_actuales)
          ].copy()
          return pd.concat(
              [df_nuevas, hist_pendientes], ignore_index=True
          ).drop_duplicates(subset=["op_clean"])
        except Exception:
          return df_nuevas

      df_tarjetas_final = consolidar_tarjetas_historico(
          df_tarj_nuevas, pend_ant_file
      )

      wb_out = openpyxl.Workbook()
      wb_out.remove(wb_out.active)

      navy_fill = PatternFill(
          start_color="000080", end_color="000080", fill_type="solid"
      )
      header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
      title_font = Font(name="Calibri", size=13, bold=True, color="000000")
      data_font = Font(name="Calibri", size=10)
      thin_border = Border(
          left=Side(style="thin", color="DDDDDD"),
          right=Side(style="thin", color="DDDDDD"),
          top=Side(style="thin", color="DDDDDD"),
          bottom=Side(style="thin", color="DDDDDD"),
      )

      def build_bank_sheet(wb, title_sheet, main_title, df_data):
        ws = wb.create_sheet(title=title_sheet)
        ws.views.sheetView[0].showGridLines = True
        ws.cell(row=2, column=2, value=main_title).font = title_font

        headers = [
            "Fecha",
            "Descripción operación",
            "Monto",
            "Nº operación",
            "Usuario",
            "Detalle de la Operación - Rspta BCP",
            "RUC / DNI",
        ]
        for c_idx, h in enumerate(headers, start=2):
          cell = ws.cell(row=4, column=c_idx, value=h)
          cell.fill = navy_fill
          cell.font = header_font
          cell.alignment = Alignment(horizontal="center", vertical="center")

        r_idx = 5
        df_sorted = df_data.copy()
        df_sorted["Fecha_sort"] = df_sorted["Fecha"].apply(smart_parse_date)
        df_sorted = df_sorted.sort_values(by="Fecha_sort", ascending=False)

        for _, row in df_sorted.iterrows():
          f_val = (
              row["Fecha_sort"].strftime("%Y-%m-%d")
              if pd.notna(row["Fecha_sort"])
              else str(row.get("Fecha", ""))[:10]
          )
          ws.cell(row=r_idx, column=2, value=f_val).alignment = Alignment(
              horizontal="center"
          )
          ws.cell(
              row=r_idx, column=3, value=str(row.get("Descripción operación", ""))
          )

          m_cell = ws.cell(
              row=r_idx, column=4, value=safe_float(row.get("Monto", 0))
          )
          m_cell.number_format = "#,##0.00"
          m_cell.alignment = Alignment(horizontal="right")

          c_op = ws.cell(
              row=r_idx, column=5, value=str(row.get("Nº operación", ""))
          )
          c_op.number_format = "@"
          c_op.alignment = Alignment(horizontal="center")

          ws.cell(
              row=r_idx,
              column=6,
              value=str(
                  row.get("Usuario", "") if pd.notna(row.get("Usuario")) else ""
              ),
          ).alignment = Alignment(horizontal="center")
          ws.cell(
              row=r_idx,
              column=7,
              value=str(
                  row.get("Detalle de la Operación - Rspta BCP", "")
                  if pd.notna(row.get("Detalle de la Operación - Rspta BCP"))
                  else ""
              ),
          )
          ws.cell(
              row=r_idx,
              column=8,
              value=str(
                  row.get("RUC / DNI", "")
                  if pd.notna(row.get("RUC / DNI"))
                  else ""
              ),
          ).alignment = Alignment(horizontal="center")

          for c in range(2, 9):
            cell = ws.cell(row=r_idx, column=c)
            cell.font = data_font
            cell.border = thin_border
          r_idx += 1

        ws.column_dimensions["A"].width = 3.0
        widths = {
            "B": 14,
            "C": 28,
            "D": 14,
            "E": 16,
            "F": 12,
            "G": 38,
            "H": 18,
        }
        for col_l, w in widths.items():
          ws.column_dimensions[col_l].width = w

      build_bank_sheet(
          wb_out,
          "FP Soles",
          "DEPOSITOS NO IDENTIFICADOS - SOLES",
          soles_final,
      )
      build_bank_sheet(
          wb_out,
          "FP Dolares",
          "DEPOSITOS NO IDENTIFICADOS - DOLARES",
          dolares_final,
      )

      if not df_tarjetas_final.empty:
        ws_t = wb_out.create_sheet(title="TARJETAS")
        ws_t.views.sheetView[0].showGridLines = True
        ws_t.cell(
            row=2, column=2, value="TARJETAS NO IDENTIFICADOS"
        ).font = title_font

        headers_t = [
            "RUC",
            "Cod. Comercio",
            "Moneda",
            "Fecha operación",
            "ID Operación",
            "Monto",
            "Cliente",
        ]
        for c_idx, h in enumerate(headers_t, start=2):
          cell = ws_t.cell(row=4, column=c_idx, value=h)
          cell.fill = navy_fill
          cell.font = header_font
          cell.alignment = Alignment(horizontal="center", vertical="center")

        r_idx = 5
        df_t_sorted = df_tarjetas_final.copy()
        df_t_sorted["Fecha_sort"] = df_t_sorted["Fecha operación"].apply(
            smart_parse_date
        )
        df_t_sorted = df_t_sorted.sort_values(by="Fecha_sort", ascending=False)

        for _, row in df_t_sorted.iterrows():
          ws_t.cell(
              row=r_idx,
              column=2,
              value=str(row.get("RUC", "20603632479")),
          ).alignment = Alignment(horizontal="center")
          ws_t.cell(
              row=r_idx, column=3, value=str(row.get("Cod. Comercio", ""))
          ).alignment = Alignment(horizontal="center")
          ws_t.cell(
              row=r_idx, column=4, value=str(row.get("Moneda", ""))
          ).alignment = Alignment(horizontal="center")

          f_op = (
              row["Fecha_sort"].strftime("%Y-%m-%d")
              if pd.notna(row["Fecha_sort"])
              else str(row.get("Fecha operación", ""))[:10]
          )
          ws_t.cell(row=r_idx, column=5, value=f_op).alignment = Alignment(
              horizontal="center"
          )

          c_op_t = ws_t.cell(
              row=r_idx, column=6, value=str(row.get("ID Operación", ""))
          )
          c_op_t.number_format = "@"
          c_op_t.alignment = Alignment(horizontal="center")

          m_cell = ws_t.cell(
              row=r_idx, column=7, value=safe_float(row.get("Monto", 0))
          )
          m_cell.number_format = "#,##0.00"
          m_cell.alignment = Alignment(horizontal="right")

          ws_t.cell(row=r_idx, column=8, value=str(row.get("Cliente", "")))

          for c in range(2, 9):
            cell = ws_t.cell(row=r_idx, column=c)
            cell.font = data_font
            cell.border = thin_border
          r_idx += 1

        ws_t.column_dimensions["A"].width = 3.0
        widths_t = {
            "B": 16,
            "C": 15,
            "D": 12,
            "E": 16,
            "F": 22,
            "G": 14,
            "H": 35,
        }
        for col_l, w in widths_t.items():
          ws_t.column_dimensions[col_l].width = w

      output = io.BytesIO()
      wb_out.save(output)
      output.seek(0)

      filename = f"Pendientes al {datetime.now().strftime('%d%m%Y')}.xlsx"
      st.success(f"Reporte '{filename}' procesado con éxito.")
      st.download_button(
          label="📥 Descargar Archivo Excel",
          data=output,
          file_name=filename,
          mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
      )