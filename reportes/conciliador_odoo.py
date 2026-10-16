import io
import re
from datetime import datetime
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import pandas as pd
import streamlit as st


def clean_op(val):
  if pd.isna(val) or val is None:
    return ""
  val_str = str(val).strip().replace(".0", "")
  val_clean = re.sub(r"[^a-zA-Z0-9]", "", val_str)
  return val_clean.lstrip("0") if val_clean.lstrip("0") else val_clean


def format_op_display(val):
  if pd.isna(val) or val is None:
    return ""
  s = str(val).strip()
  return s[:-2] if s.endswith(".0") else s


def safe_float(val):
  if pd.isna(val) or val is None:
    return 0.0
  try:
    return float(str(val).replace(",", "").strip())
  except Exception:
    return 0.0


def smart_parse_date(val):
  if pd.isna(val) or val is None:
    return pd.NaT
  if isinstance(val, (datetime, pd.Timestamp)):
    return pd.to_datetime(val)
  s = str(val).strip()
  if re.match(r"^\d{4}[-/]\d{1,2}[-/]\d{1,2}", s):
    return pd.to_datetime(s, format="mixed", errors="coerce")
  return pd.to_datetime(s, dayfirst=True, errors="coerce")


def load_smart(file_obj, sheet_name=0, keywords=["fecha", "monto"]):
  df_raw = pd.read_excel(file_obj, sheet_name=sheet_name, header=None)
  header_idx = 0
  for idx, row in df_raw.head(15).iterrows():
    row_str = " ".join([str(v).lower() for v in row.values])
    if sum(1 for kw in keywords if kw in row_str) >= len(keywords):
      header_idx = idx
      break
  file_obj.seek(0)
  df = pd.read_excel(
      file_obj, sheet_name=sheet_name, header=header_idx, dtype=str
  )
  df.columns = [str(c).strip() for c in df.columns]
  return df


def render():
  st.subheader("Conciliador Semanal: Odoo vs Bancos")

  uploaded_files = st.file_uploader(
      "Sube los archivos necesarios (Odoo, BCP, Tarjetas, Bancos Anterior):",
      type=["xlsx"],
      accept_multiple_files=True,
      key="uploader_conciliador",
  )

  if st.button("Generar Reporte Conciliador", type="primary"):
    if not uploaded_files:
      st.error("Debes subir los archivos requeridos.")
      return

    file_dict = {f.name: f for f in uploaded_files}

    def find_file(keywords):
      for name, f in file_dict.items():
        if all(k.lower() in name.lower() for k in keywords):
          return f
      return None

    odoo_file = find_file(["odoo"])
    bcp_file = (
        find_file(["bcp", "movimiento"])
        or find_file(["bcp", "diario"])
        or find_file(["bcp"])
    )
    tarj_file = find_file(["tarjeta"])

    bancos_ant_file = None
    for name, f in file_dict.items():
      n_l = name.lower()
      if (
          "banco" in n_l
          and f != bcp_file
          and "odoo" not in n_l
          and "tarjeta" not in n_l
      ):
        bancos_ant_file = f
        break

    if not (odoo_file and bcp_file):
      st.error("Es obligatorio subir tanto el archivo de Odoo como el de BCP.")
      return

    with st.spinner("Procesando cruce de conciliación..."):
      df_odoo = load_smart(
          odoo_file, sheet_name=0, keywords=["fecha", "diario", "importe"]
      )
      estado_col = [c for c in df_odoo.columns if "estado" in c.lower()]
      if estado_col:
        df_odoo = df_odoo[
            df_odoo[estado_col[0]].astype(str).str.strip().str.lower()
            != "cancelado"
        ].copy()

      df_odoo["op_display"] = df_odoo["Numero de operacion"].apply(
          format_op_display
      )
      df_odoo["op_clean"] = df_odoo["Numero de operacion"].apply(clean_op)
      df_odoo["Fecha_dt"] = df_odoo["Fecha"].apply(smart_parse_date)
      df_odoo["Fecha_op_dt"] = df_odoo["Fecha de operacion"].apply(
          smart_parse_date
      )
      df_odoo["Origen"] = "ODOO_ACTUAL"

      max_fecha_corte = df_odoo["Fecha_dt"].dropna().max()
      mes_corte = max_fecha_corte.month
      anio_corte = max_fecha_corte.year
      fecha_corte_str = max_fecha_corte.strftime("%d %m %y")

      xls_bcp = pd.ExcelFile(bcp_file)
      sheet_bcp = (
          "Movimientos"
          if "Movimientos" in xls_bcp.sheet_names
          else xls_bcp.sheet_names[0]
      )
      df_bcp = load_smart(
          bcp_file, sheet_name=sheet_bcp, keywords=["fecha", "monto"]
      )
      df_bcp["Fecha_dt"] = df_bcp["Fecha"].apply(smart_parse_date)
      df_bcp["op_display"] = df_bcp["Nº operación"].apply(format_op_display)
      df_bcp["op_clean"] = df_bcp["Nº operación"].apply(clean_op)
      df_bcp["Monto_num"] = df_bcp["Monto"].apply(safe_float)

      is_egreso = (
          df_bcp["OBSERVACIONES"]
          .astype(str)
          .str.contains("egreso", case=False, na=False)
      )
      df_bcp_mes = df_bcp[
          (df_bcp["Fecha_dt"].dt.month == mes_corte)
          & (df_bcp["Fecha_dt"].dt.year == anio_corte)
          & (df_bcp["Monto_num"] > 0)
          & (~is_egreso)
      ].copy()

      mapa_tarjeta_a_banco = {}
      mapa_banco_display = {}
      tarjetas_detalle_no_odoo = []

      if tarj_file:
        wb_tarj = openpyxl.load_workbook(tarj_file, data_only=True)
        for s in [
            "Culqi POS MN",
            "Culqi POS ME",
            "Culqi Link MN",
            "Culqi Link ME",
        ]:
          if s in wb_tarj.sheetnames:
            ws_t = wb_tarj[s]
            for r in range(5, ws_t.max_row + 1):
              val_i = ws_t.cell(row=r, column=9).value
              val_z = ws_t.cell(row=r, column=26).value
              raw_ah = ws_t.cell(row=r, column=34).value
              raw_bc = ws_t.cell(row=r, column=55).value

              val_ah, val_bc = clean_op(raw_ah), clean_op(raw_bc)
              disp_ah, disp_bc = format_op_display(raw_ah), format_op_display(
                  raw_bc
              )

              if val_ah and val_bc:
                mapa_tarjeta_a_banco[val_ah] = val_bc
                mapa_banco_display[val_bc] = disp_bc
                f_dt = smart_parse_date(val_i)
                if (
                    pd.notna(f_dt)
                    and f_dt.month == mes_corte
                    and f_dt.year == anio_corte
                ):
                  tarjetas_detalle_no_odoo.append({
                      "op_clean": val_ah,
                      "op_display": disp_ah,
                      "Numero de operacion": disp_ah,
                      "Importe con signo": safe_float(val_z),
                      "Fecha de operacion": str(val_i)[:10],
                      "op_banco_puente": val_bc,
                      "Origen": "TARJETA_NO_ODOO",
                  })

        if "MPago" in wb_tarj.sheetnames:
          ws_mp = wb_tarj["MPago"]
          current_cobros = []
          for r in range(2, ws_mp.max_row + 1):
            tipo = str(ws_mp.cell(row=r, column=2).value).strip()
            raw_rel = ws_mp.cell(row=r, column=4).value
            raw_banco = ws_mp.cell(row=r, column=8).value
            rel, banco_op = clean_op(raw_rel), clean_op(raw_banco)
            disp_banco = format_op_display(raw_banco)

            if tipo == "Retiro de dinero a cuenta bancaria":
              if banco_op:
                for cobro_rel in current_cobros:
                  mapa_tarjeta_a_banco[cobro_rel] = banco_op
                  mapa_banco_display[banco_op] = disp_banco
              current_cobros = []
            elif rel:
              current_cobros.append(rel)

      df_odoo["op_banco_puente"] = (
          df_odoo["op_clean"].map(mapa_tarjeta_a_banco).fillna("")
      )

      odoo_transito_anterior = []
      if bancos_ant_file:
        try:
          wb_ba = openpyxl.load_workbook(bancos_ant_file, data_only=True)
          ws_ba = wb_ba.active
          for r in range(2, ws_ba.max_row + 1):
            c_fill = ws_ba.cell(row=r, column=1).fill
            is_celeste = bool(
                c_fill
                and c_fill.fgColor
                and getattr(c_fill.fgColor, "theme", None) == 4
            )
            f_val = ws_ba.cell(row=r, column=1).value
            f_dt = smart_parse_date(f_val)
            if is_celeste or (
                pd.notna(f_dt)
                and (f_dt.month < mes_corte or f_dt.year < anio_corte)
            ):
              raw_op = ws_ba.cell(row=r, column=6).value
              raw_b_op = ws_ba.cell(row=r, column=15).value
              op_val, b_op_val = clean_op(raw_op), clean_op(raw_b_op)
              if op_val:
                odoo_transito_anterior.append({
                    "Fecha": ws_ba.cell(row=r, column=1).value,
                    "Nombre": ws_ba.cell(row=r, column=2).value,
                    "Diario": ws_ba.cell(row=r, column=3).value,
                    "Cliente/Proveedor": ws_ba.cell(row=r, column=4).value,
                    "Importe con signo": ws_ba.cell(row=r, column=5).value,
                    "Numero de operacion": format_op_display(raw_op),
                    "op_display": format_op_display(raw_op),
                    "Fecha de operacion": ws_ba.cell(row=r, column=7).value,
                    "Descripción operación": ws_ba.cell(
                        row=r, column=8
                    ).value,
                    "Estado pago": ws_ba.cell(row=r, column=9).value,
                    "Facturas conciliadas": ws_ba.cell(
                        row=r, column=10
                    ).value,
                    "op_clean": op_val,
                    "Fecha_dt": f_dt,
                    "Fecha_op_dt": smart_parse_date(
                        ws_ba.cell(row=r, column=7).value
                    ),
                    "op_banco_puente": b_op_val
                    or mapa_tarjeta_a_banco.get(op_val, ""),
                    "Origen": "BANCOS_ANTERIOR_TRANSITO",
                })
        except Exception as e:
          st.warning(f"Aviso bancos anterior: {e}")

      if odoo_transito_anterior:
        df_ant = pd.DataFrame(odoo_transito_anterior)
        ops_actuales = set(df_odoo["op_clean"])
        df_ant = df_ant[~df_ant["op_clean"].isin(ops_actuales)].copy()
        df_odoo = pd.concat([df_odoo, df_ant], ignore_index=True)

      ops_odoo_reg = set(df_odoo[df_odoo["op_clean"] != ""]["op_clean"])
      tarj_sin_odoo = [
          t
          for t in tarjetas_detalle_no_odoo
          if t["op_clean"] not in ops_odoo_reg
      ]
      if tarj_sin_odoo:
        df_tarj_no_odoo = pd.DataFrame(tarj_sin_odoo).drop_duplicates(
            subset=["op_clean"]
        )
        for c in [
            "Fecha",
            "Nombre",
            "Diario",
            "Cliente/Proveedor",
            "Descripción operación",
            "Estado pago",
            "Facturas conciliadas",
            "Fecha_dt",
            "Fecha_op_dt",
        ]:
          df_tarj_no_odoo[c] = None
        df_tarj_no_odoo["Fecha_op_dt"] = df_tarj_no_odoo[
            "Fecha de operacion"
        ].apply(smart_parse_date)
        df_odoo = pd.concat([df_odoo, df_tarj_no_odoo], ignore_index=True)

      obs_b = df_bcp_mes["OBSERVACIONES"].astype(str).str.strip()
      is_especial_banco = obs_b.str.contains(
          "Cambio de Moneda|DUPLICADO BCP|Transferencia No Realizada|Extorno",
          case=False,
          na=False,
      )
      is_tarjeta_banco = obs_b.str.contains(
          "OK Tarjetas", case=False, na=False
      ) | df_bcp_mes["Descripción operación"].astype(str).str.contains(
          "CULQI|INCUBADORA", case=False, na=False
      )

      df_bcp_especiales = df_bcp_mes[is_especial_banco].copy()
      df_bcp_tarjetas = df_bcp_mes[
          is_tarjeta_banco & (~is_especial_banco)
      ].copy()
      df_bcp_reg = df_bcp_mes[(~is_especial_banco) & (~is_tarjeta_banco)].copy()

      desc_o = df_odoo["Descripción operación"].astype(str).str.upper()
      is_tarjeta_odoo = desc_o.str.contains(
          "FERRETTI|MERCADO PAGO", case=False, na=False
      ) | (df_odoo["op_banco_puente"] != "")
      is_mes_ant = (df_odoo["Fecha_op_dt"].dt.month < mes_corte) | (
          df_odoo["Fecha_op_dt"].dt.year < anio_corte
      )

      df_odoo_meses_ant = df_odoo[
          (~is_tarjeta_odoo)
          & is_mes_ant
          & (df_odoo["Origen"] != "TARJETA_NO_ODOO")
      ].copy()
      df_odoo_reg = df_odoo[
          (~is_tarjeta_odoo)
          & (~is_mes_ant)
          & (df_odoo["Origen"] != "TARJETA_NO_ODOO")
      ].copy()
      df_odoo_tarj = df_odoo[is_tarjeta_odoo].copy()

      cruce_1a1 = df_bcp_reg.merge(
          df_odoo_reg, on="op_clean", how="inner", suffixes=("_bcp", "_odoo")
      ).drop_duplicates(subset=["op_clean"])
      sort_col = (
          "Monto_num_bcp" if "Monto_num_bcp" in cruce_1a1.columns else "Monto_num"
      )
      cruce_1a1 = cruce_1a1.sort_values(by=sort_col, ascending=True)

      ops_matched_1a1 = set(cruce_1a1["op_clean"])
      df_bcp_especiales_sorted = df_bcp_especiales.sort_values(
          by=["OBSERVACIONES", "Fecha_dt"], ascending=[True, True]
      )
      df_bcp_no_id = df_bcp_reg[
          ~df_bcp_reg["op_clean"].isin(ops_matched_1a1)
      ].copy()
      df_bcp_no_id_sorted = df_bcp_no_id.sort_values(
          by=["Fecha_dt", "Monto_num"], ascending=[True, True]
      )

      bcp_tarj_ops = set(df_bcp_tarjetas["op_clean"])
      df_tarj_conciliadas = df_odoo_tarj[
          df_odoo_tarj["op_banco_puente"].isin(bcp_tarj_ops)
      ].copy()
      mapa_fecha_banco = df_bcp_tarjetas.set_index("op_clean")[
          "Fecha_dt"
      ].to_dict()
      df_tarj_conciliadas["Fecha_banco_lote"] = df_tarj_conciliadas[
          "op_banco_puente"
      ].map(mapa_fecha_banco)
      df_tarj_conciliadas = df_tarj_conciliadas.sort_values(
          by=["Fecha_banco_lote", "op_banco_puente", "Fecha_dt"],
          ascending=[True, True, True],
      )

      df_tarj_sin_conciliar = df_odoo_tarj[
          ~df_odoo_tarj["op_banco_puente"].isin(bcp_tarj_ops)
      ].sort_values(
          by=["Descripción operación", "Diario", "Fecha_dt"],
          ascending=[True, True, True],
      )

      wb = openpyxl.Workbook()
      ws = wb.active
      ws.title = "Sheet1"
      ws.views.sheetView[0].showGridLines = True

      fill_banco_header = PatternFill(
          start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"
      )
      fill_yellow = PatternFill(
          start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"
      )
      fill_red_light = PatternFill(
          start_color="F8CBAD", end_color="F8CBAD", fill_type="solid"
      )
      fill_orange_light = PatternFill(
          start_color="F8CBAD", end_color="F8CBAD", fill_type="solid"
      )
      fill_blue_light = PatternFill(
          start_color="DDEBF7", end_color="DDEBF7", fill_type="solid"
      )
      fill_cust_out = PatternFill(
          start_color="F2DCDB", end_color="F2DCDB", fill_type="solid"
      )

      font_header = Font(name="Calibri", size=10, bold=True)
      font_data = Font(name="Calibri", size=10)
      font_cust_out = Font(name="Calibri", size=10, color="58151C", bold=True)
      font_red_tarjeta = Font(
          name="Calibri", size=10, color="FF0000", bold=True
      )
      thin_border = Border(
          left=Side(style="thin", color="E0E0E0"),
          right=Side(style="thin", color="E0E0E0"),
          top=Side(style="thin", color="E0E0E0"),
          bottom=Side(style="thin", color="E0E0E0"),
      )

      headers_o = [
          "Fecha",
          "Nombre",
          "Diario",
          "Cliente/Proveedor",
          "Importe con signo",
          "Numero de operacion",
          "Fecha de operacion",
          "Descripción operación",
          "Estado pago",
          "Facturas conciliadas",
      ]
      headers_b = [
          "Fecha",
          "Descripción operación",
          "Monto",
          "Nº operación",
          "Moneda",
          "OBSERVACIONES",
      ]

      for c, h in enumerate(headers_o, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center")

      for c, h in enumerate(headers_b, start=12):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = font_header
        cell.fill = fill_banco_header
        cell.alignment = Alignment(horizontal="center")

      def write_odoo_row(ws, r, row, fill=None):
        if row is None or (isinstance(row, pd.Series) and row.empty):
          return
        nombre_val = str(
            row.get("Nombre", "") if pd.notna(row.get("Nombre")) else ""
        )
        origen = row.get("Origen", "")
        is_cust_out = nombre_val.startswith("CUST.OUT")
        is_no_odoo = origen == "TARJETA_NO_ODOO"

        row_fill = fill_cust_out if is_cust_out else fill
        row_font = (
            font_red_tarjeta
            if is_no_odoo
            else (font_cust_out if is_cust_out else font_data)
        )

        f_val = (
            str(row.get("Fecha", ""))[:10]
            if pd.notna(row.get("Fecha")) and not is_no_odoo
            else ""
        )
        ws.cell(row=r, column=1, value=f_val).alignment = Alignment(
            horizontal="center"
        )
        ws.cell(
            row=r, column=2, value=nombre_val if not is_no_odoo else ""
        ).alignment = Alignment(horizontal="left")
        ws.cell(
            row=r,
            column=3,
            value=(
                str(row.get("Diario", ""))
                if not is_no_odoo and pd.notna(row.get("Diario"))
                else ""
            ),
        )
        ws.cell(
            row=r,
            column=4,
            value=(
                str(row.get("Cliente/Proveedor", ""))
                if not is_no_odoo and pd.notna(row.get("Cliente/Proveedor"))
                else ""
            ),
        )

        m_cell = ws.cell(
            row=r, column=5, value=safe_float(row.get("Importe con signo", 0))
        )
        m_cell.number_format = "#,##0.00"
        m_cell.alignment = Alignment(horizontal="right")

        op_f = str(
            row.get("op_display", "")
            if pd.notna(row.get("op_display"))
            else row.get("Numero de operacion", "")
        )
        c_op_f = ws.cell(row=r, column=6, value=op_f)
        c_op_f.number_format = "@"
        c_op_f.alignment = Alignment(horizontal="center")

        ws.cell(
            row=r,
            column=7,
            value=(
                str(row.get("Fecha de operacion", ""))[:10]
                if pd.notna(row.get("Fecha de operacion"))
                else ""
            ),
        ).alignment = Alignment(horizontal="center")
        ws.cell(
            row=r,
            column=8,
            value=(
                str(row.get("Descripción operación", ""))
                if not is_no_odoo and pd.notna(row.get("Descripción operación"))
                else ""
            ),
        )
        ws.cell(
            row=r,
            column=9,
            value=(
                str(row.get("Estado pago", ""))
                if not is_no_odoo and pd.notna(row.get("Estado pago"))
                else ""
            ),
        ).alignment = Alignment(horizontal="center")
        ws.cell(
            row=r,
            column=10,
            value=(
                str(row.get("Facturas conciliadas", ""))
                if not is_no_odoo and pd.notna(row.get("Facturas conciliadas"))
                else ""
            ),
        ).alignment = Alignment(horizontal="center")

        for c in range(1, 11):
          c_item = ws.cell(row=r, column=c)
          c_item.font = row_font
          c_item.border = thin_border
          if row_fill and not is_no_odoo:
            c_item.fill = row_fill

      def write_banco_row(ws, r, row, fill=None):
        if row is None or (isinstance(row, pd.Series) and row.empty):
          return
        f_str = (
            row["Fecha_dt"].strftime("%Y-%m-%d")
            if pd.notna(row.get("Fecha_dt"))
            else str(row.get("Fecha", ""))[:10]
        )
        ws.cell(row=r, column=12, value=f_str).alignment = Alignment(
            horizontal="center"
        )
        ws.cell(row=r, column=13, value=str(row.get("Descripción operación", "")))

        m_cell = ws.cell(
            row=r, column=14, value=safe_float(row.get("Monto", 0))
        )
        m_cell.number_format = "#,##0.00"
        m_cell.alignment = Alignment(horizontal="right")

        op_o = str(
            row.get("op_display", "")
            if pd.notna(row.get("op_display"))
            else row.get("Nº operación", "")
        )
        c_op_o = ws.cell(row=r, column=15, value=op_o)
        c_op_o.number_format = "@"
        c_op_o.alignment = Alignment(horizontal="center")

        ws.cell(row=r, column=16, value=str(row.get("Moneda", ""))).alignment = (
            Alignment(horizontal="center")
        )
        ws.cell(
            row=r,
            column=17,
            value=str(
                row.get("OBSERVACIONES", "")
                if pd.notna(row.get("OBSERVACIONES"))
                else ""
            ),
        ).alignment = Alignment(horizontal="left")

        for c in range(12, 18):
          c_item = ws.cell(row=r, column=c)
          c_item.font = font_data
          c_item.border = thin_border
          if fill:
            c_item.fill = fill

      curr_r = 2
      banco_escritos = set()

      for _, r_b in df_bcp_especiales_sorted.iterrows():
        write_banco_row(ws, curr_r, r_b)
        banco_escritos.add(r_b["op_clean"])
        curr_r += 1

      for _, r_b in df_bcp_no_id_sorted.iterrows():
        write_banco_row(ws, curr_r, r_b)
        banco_escritos.add(r_b["op_clean"])
        curr_r += 1

      for _, r_c in cruce_1a1.iterrows():
        b_row = df_bcp_reg[df_bcp_reg["op_clean"] == r_c["op_clean"]].iloc[0]
        o_row = df_odoo_reg[df_odoo_reg["op_clean"] == r_c["op_clean"]].iloc[0]
        write_odoo_row(ws, curr_r, o_row)
        write_banco_row(ws, curr_r, b_row, fill=fill_yellow)
        banco_escritos.add(b_row["op_clean"])
        curr_r += 1

      for _, r_o in df_odoo_meses_ant.sort_values(by="Fecha_op_dt").iterrows():
        write_odoo_row(ws, curr_r, r_o, fill=fill_red_light)
        curr_r += 1

      df_odoo_reg_sin_match = df_odoo_reg[
          ~df_odoo_reg["op_clean"].isin(ops_matched_1a1)
      ].copy()
      for _, r_o in df_odoo_reg_sin_match.iterrows():
        write_odoo_row(ws, curr_r, r_o)
        curr_r += 1

      grupos_tarj_conciliadas = df_tarj_conciliadas.groupby(
          "op_banco_puente", sort=False
      )
      for op_lote, group_odoo in grupos_tarj_conciliadas:
        match_banco_lote = df_bcp_tarjetas[
            df_bcp_tarjetas["op_clean"] == op_lote
        ]
        banco_row_to_write = (
            match_banco_lote.iloc[0] if not match_banco_lote.empty else None
        )
        first_item = True
        for _, r_o in group_odoo.iterrows():
          f_dt = r_o.get("Fecha_dt")
          f_op = r_o.get("Fecha_op_dt")
          origen = r_o.get("Origen", "")

          if origen == "BANCOS_ANTERIOR_TRANSITO":
            c_fill = fill_blue_light
          elif pd.notna(f_op) and (
              (f_op.month < mes_corte) or (f_op.year < anio_corte)
          ):
            c_fill = (
                fill_orange_light
                if pd.notna(f_dt) and f_dt.month == mes_corte
                else fill_blue_light
            )
          else:
            c_fill = fill_yellow

          write_odoo_row(ws, curr_r, r_o, fill=c_fill)
          if first_item and banco_row_to_write is not None:
            write_banco_row(ws, curr_r, banco_row_to_write, fill=fill_yellow)
            banco_escritos.add(banco_row_to_write["op_clean"])
            first_item = False
          curr_r += 1

      tarjetas_restantes = df_bcp_tarjetas[
          ~df_bcp_tarjetas["op_clean"].isin(banco_escritos)
      ].sort_values(by="Fecha_dt", ascending=True)
      for _, r_b in tarjetas_restantes.iterrows():
        write_banco_row(ws, curr_r, r_b, fill=fill_yellow)
        banco_escritos.add(r_b["op_clean"])
        curr_r += 1

      banco_faltante = df_bcp_mes[
          ~df_bcp_mes["op_clean"].isin(banco_escritos)
      ].sort_values(by="Fecha_dt", ascending=True)
      for _, r_b in banco_faltante.iterrows():
        write_banco_row(ws, curr_r, r_b)
        banco_escritos.add(r_b["op_clean"])
        curr_r += 1

      for _, r_o in df_tarj_sin_conciliar.iterrows():
        write_odoo_row(ws, curr_r, r_o)
        curr_r += 1

      widths = [
          12,
          24,
          22,
          34,
          14,
          16,
          14,
          25,
          12,
          16,
          4,
          12,
          25,
          14,
          16,
          12,
          28,
      ]
      for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = w

      output = io.BytesIO()
      wb.save(output)
      output.seek(0)

      filename = f"Bancos {fecha_corte_str}.xlsx"
      st.success(f"Reporte '{filename}' procesado con éxito.")
      st.download_button(
          label="📥 Descargar Archivo Excel",
          data=output,
          file_name=filename,
          mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
      )